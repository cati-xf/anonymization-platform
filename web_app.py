import argparse
import contextlib
import io
import json
import os
import queue
import shutil
import subprocess
import sys
import threading
import traceback
import uuid
from pathlib import Path

import uvicorn
from fastapi import FastAPI, Form, HTTPException, UploadFile
from fastapi.responses import FileResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles

sys.path.insert(0, str(Path(__file__).parent))

app = FastAPI(title="Speech Anonymizer")

APP_DIR = Path(__file__).resolve().parent
UPLOADS_DIR = APP_DIR / "data" / "uploads"
OUTPUTS_DIR = APP_DIR / "data" / "outputs"
UPLOADS_DIR.mkdir(parents=True, exist_ok=True)
OUTPUTS_DIR.mkdir(parents=True, exist_ok=True)

# Local LLM (Gemma via Ollama) as an alternative to Wikidata for "gen"/"rand"
# replacement text -- see text_anon/llm/llm_combined.py. Override for a
# different locally-pulled model or a non-default Ollama host.
TEXT_LLM_MODEL = os.environ.get("TEXT_LLM_MODEL", "gemma3:12b")
TEXT_LLM_OLLAMA_HOST = os.environ.get("TEXT_LLM_OLLAMA_HOST", "http://localhost:11434")

# job registry (status, log queue, apply queue, result, error)
JOBS: dict[str, dict] = {}


def _ensure_ffmpeg_on_path() -> None:
    """Same as audio._paths.ensure_ffmpeg_on_path (not imported from there:
    importing the `audio` package pulls in whisperx/torch, too slow at startup).
    Puts a fallback ffmpeg on PATH so every library and subprocess finds it."""
    if shutil.which("ffmpeg"):
        return
    exe = os.environ.get("FFMPEG_BIN")
    if not exe:
        try:
            import imageio_ffmpeg
            exe = imageio_ffmpeg.get_ffmpeg_exe()
        except Exception:
            return
    bin_dir = APP_DIR / "data" / "bin"
    bin_dir.mkdir(parents=True, exist_ok=True)
    link = bin_dir / "ffmpeg"
    if link.is_symlink() or link.exists():
        link.unlink()
    link.symlink_to(exe)
    os.environ["PATH"] = str(bin_dir) + os.pathsep + os.environ.get("PATH", "")


_ensure_ffmpeg_on_path()


def _convert_to_wav16k(src_path: str, dst_path: str) -> None:
    """Convert any ffmpeg-readable audio (mp3, m4a, ogg, flac, webm, ...) to 16kHz mono PCM16 WAV."""
    result = subprocess.run(
        ["ffmpeg", "-y", "-i", src_path, "-ar", "16000", "-ac", "1", dst_path],
        capture_output=True,
    )
    if result.returncode != 0:
        raise HTTPException(
            status_code=400,
            detail=f"Could not read audio file: {result.stderr.decode(errors='replace')[-500:]}",
        )


# ── Logging ───────────────────────────────────────────────────────────────────

class _QueueWriter(io.TextIOBase):
    """ instead of printing to terminal redirects into queue """
    def __init__(self, q: queue.Queue):
        self._q = q

    def write(self, s: str) -> int:
        if s and s.strip():
            self._q.put({"type": "log", "msg": s.rstrip()})
        return len(s)

    def flush(self) -> None:
        pass


# ── Shared helpers ────────────────────────────────────────────────────────────

def _make_config(per: str, loc: str, org: str, misc: str, regex_method: str, whitelist: str = "", regex_skip: str = ""):
    import re as _re
    from text_anon import AnonymizationConfig
    wl = [w.strip() for w in _re.split(r"[,\n]+", whitelist) if w.strip()]
    skip = [s.strip() for s in _re.split(r"[,\n]+", regex_skip) if s.strip()]
    return AnonymizationConfig(
        regex=regex_method, per=per, loc=loc, org=org, misc=misc,
        whitelist=wl, regex_skip_labels=skip,
    )


def _detect_entities(text: str, config, replacement_source: str, local_only: bool = False) -> list[dict]:
    """Detect entities, sourcing "gen"/"rand" replacement text from either
    Wikidata (default, unchanged baseline) or a locally-hosted LLM (Gemma via
    Ollama) per `replacement_source` ("wikidata"/"llm"). sup/tag/none labels
    and regex detection are identical either way -- neither needs an
    external source. Raises RuntimeError with a clear message if the LLM is
    unreachable, so callers don't have to interpret a raw requests.* traceback."""
    if replacement_source == "llm":
        from text_anon.llm.llm_combined import detect_entities_llm_combined
        try:
            return detect_entities_llm_combined(text, config, model=TEXT_LLM_MODEL, ollama_host=TEXT_LLM_OLLAMA_HOST)
        except Exception as e:
            raise RuntimeError(
                f"Could not reach the local LLM (Ollama at {TEXT_LLM_OLLAMA_HOST}) — "
                f"make sure Ollama is running and '{TEXT_LLM_MODEL}' is pulled. ({e})"
            ) from e
    from text_anon import detect_entities
    return detect_entities(text, config, local_only=local_only)


def _cuda_available() -> bool:
    try:
        import torch
        return torch.cuda.is_available()
    except ImportError:
        return False


def _release_gpu() -> None:
    import gc
    gc.collect()
    try:
        import torch
        if torch.cuda.is_available():
            torch.cuda.empty_cache()
    except ImportError:
        pass


def _knn_ref_wavs(ref_set: str) -> list[str]:
    """kNN-VC target-speaker clips (ref/tmp/<ref_set>*.wav). They are not in the
    git repository, so raise a clear error instead of letting kNN-VC crash on an
    empty list (or, in full_infilling_knn, silently skip the voice conversion)."""
    ref_dir = Path(__file__).parent / "ref" / "tmp"
    wavs = sorted(str(p) for p in ref_dir.glob(f"{ref_set}*.wav"))
    if not wavs:
        raise RuntimeError(
            f"No kNN-VC reference clips for speaker '{ref_set}' in ref/tmp/. "
            "They are not included in the repository; see README step 6."
        )
    return wavs


def _make_job(noise_mode: str = "remove", batch_mode: bool = False) -> dict:
    log_q: queue.Queue = queue.Queue()
    apply_q: queue.Queue = queue.Queue(maxsize=1)
    return {
        "status": "running",
        "log_q": log_q,
        "apply_q": apply_q,
        "result": None,
        "error": None,
        "noise_mode": noise_mode,
        "batch_mode": batch_mode,
        "_noise_path": None,
        "_original_audio_path": None,
    }


def _separate_if_needed(job: dict, audio_path: str, log, writer, device: str) -> str:
    """Run USS separation; stores noise/original paths in job. Returns working audio path."""
    from audio.noise_separation import separate_speech_noise
    speech_path = audio_path + ".speech.wav"
    noise_path = audio_path + ".noise.wav"
    log("▶ Step 0: Separating speech from background noise (USS)")
    with contextlib.redirect_stdout(writer), contextlib.redirect_stderr(writer):
        separate_speech_noise(audio_path, speech_path, noise_path, device=device)
    job["_noise_path"] = noise_path
    job["_original_audio_path"] = audio_path
    _release_gpu()
    return speech_path


def _mix_noise_if_needed(job: dict, out_path: str, segment_map: list, log, writer) -> str:
    """Re-add stored noise into the anonymized audio; returns final output path."""
    noise_path = job.get("_noise_path")
    original_path = job.get("_original_audio_path")
    if not noise_path:
        log("  [WARN] No noise file stored, skipping noise mix")
        return out_path
    from audio.mix_noise import mix_noise as _mix_noise_fn
    mixed_out = out_path.replace(".wav", "_mixed.wav")
    log("▶ Re-adding background noise to anonymized audio")
    with contextlib.redirect_stdout(writer), contextlib.redirect_stderr(writer):
        _mix_noise_fn(out_path, noise_path, segment_map, mixed_out, original_path=original_path)
    return mixed_out


# ── Pipeline workers ──────────────────────────────────────────────────────────

def _worker_detect_infilling(
    job_id: str,
    audio_path: str,
    per: str,
    loc: str,
    org: str,
    misc: str,
    regex_method: str,
    use_knn: bool,
    ref_set: str = "",
    whitelist: str = "",
    regex_skip: str = "",
    batch_mode: bool = False,
    replacement_source: str = "wikidata",
) -> None:
    """Phase 1: ASR + entity detection → wait for user review → Phase 2: synthesis.

    In batch_mode, phase 1's auto-detected entities are applied directly with no
    human review/edit step — the job never pauses in "review" status. Used when
    the frontend submits more than one file at once.
    """
    job = JOBS[job_id]
    log_q: queue.Queue = job["log_q"]
    apply_q: queue.Queue = job["apply_q"]
    writer = _QueueWriter(log_q)
    noise_mode: str = job["noise_mode"]

    def log(msg: str) -> None:
        log_q.put({"type": "log", "msg": msg})

    working_audio = audio_path
    device = "cuda" if _cuda_available() else "cpu"

    try:
        if noise_mode in ("remove", "keep"):
            working_audio = _separate_if_needed(job, audio_path, log, writer, device)

        with contextlib.redirect_stdout(writer), contextlib.redirect_stderr(writer):
            from audio.asr import get_timed_transcript

            config = _make_config(per, loc, org, misc, regex_method, whitelist, regex_skip)

            log("▶ Step 1: Transcribing with word timestamps")
            words = get_timed_transcript(working_audio, device=device)
            transcript = " ".join(w.word for w in words)
            log(f"Transcript: {transcript[:100]}{'…' if len(transcript) > 100 else ''}")

            log("▶ Step 2: Detecting sensitive entities" + (" (local LLM)" if replacement_source == "llm" else ""))
            entities = _detect_entities(transcript, config, replacement_source)
            log(f"Found {len(entities)} entities")
            for e in entities:
                log(f"  [{e['start']}:{e['end']}] '{e['original']}' → '{e['replacement']}'")

        # Save transcript + words for phase 2
        job["_words"] = words
        job["_transcript"] = transcript

        if batch_mode:
            log("⚠ Batch mode: skipping human review — applying auto-detected entities directly")
            user_entities = entities
        else:
            job["status"] = "review"
            log_q.put({"type": "review", "transcript": transcript, "entities": entities})

            # Wait for user's reviewed entities (30-minute timeout)
            try:
                user_entities = apply_q.get(timeout=1800)
            except queue.Empty:
                job["status"] = "error"
                job["error"] = "Review session expired"
                log_q.put({"type": "error", "msg": "Review session timed out"})
                return

        job["status"] = "running"
        log("▶ Step 3: Synthesizing with " + ("auto-detected entities (batch mode)" if batch_mode else "reviewed entities"))

        with contextlib.redirect_stdout(writer), contextlib.redirect_stderr(writer):
            from audio.pipelines.infilling_f5_tts import synthesize_anonymized

            out_path = str(OUTPUTS_DIR / f"{job_id}.wav")
            ref_wavs = _knn_ref_wavs(ref_set) if use_knn else None
            anonymized_text, _, segment_map = synthesize_anonymized(
                working_audio, out_path, words, user_entities,
                ref_wav_paths=ref_wavs, device=device,
            )

        if noise_mode == "keep":
            out_path = _mix_noise_if_needed(job, out_path, segment_map, log, writer)

        job["_audio_path"] = working_audio
        job["_entities"] = user_entities
        job["_ref_wavs"] = ref_wavs
        job["result"] = {
            "transcript": transcript,
            "anonymized": anonymized_text,
            "output_audio": out_path,
            "segment_map": segment_map,
        }
        job["status"] = "done"
        log_q.put({"type": "done"})

    except Exception as e:
        job["status"] = "error"
        job["error"] = str(e)
        log_q.put({"type": "log", "msg": traceback.format_exc()})
        log_q.put({"type": "error", "msg": str(e)})
    finally:
        _release_gpu()


def _worker_detect_asr_tts(
    job_id: str,
    audio_path: str,
    per: str,
    loc: str,
    org: str,
    misc: str,
    regex_method: str,
    f5_voice: str = "original",
    whitelist: str = "",
    regex_skip: str = "",
    batch_mode: bool = False,
    replacement_source: str = "wikidata",
) -> None:
    """Phase 1: ASR + entity detection → wait for user review → Phase 2: full TTS.

    In batch_mode, phase 1's auto-detected entities are applied directly with no
    human review/edit step — the job never pauses in "review" status. Used when
    the frontend submits more than one file at once.
    """
    job = JOBS[job_id]
    log_q: queue.Queue = job["log_q"]
    apply_q: queue.Queue = job["apply_q"]
    writer = _QueueWriter(log_q)
    noise_mode: str = job["noise_mode"]
    device = "cuda" if _cuda_available() else "cpu"

    def log(msg: str) -> None:
        log_q.put({"type": "log", "msg": msg})

    working_audio = audio_path

    try:
        if noise_mode in ("remove", "keep"):
            working_audio = _separate_if_needed(job, audio_path, log, writer, device)

        with contextlib.redirect_stdout(writer), contextlib.redirect_stderr(writer):
            from audio.asr import get_transcript

            config = _make_config(per, loc, org, misc, regex_method, whitelist, regex_skip)

            log("▶ Step 1: Transcribing audio")
            transcript = get_transcript(working_audio)
            log(f"Transcript: {transcript[:100]}{'…' if len(transcript) > 100 else ''}")

            log("▶ Step 2: Detecting sensitive entities" + (" (local LLM)" if replacement_source == "llm" else ""))
            entities = _detect_entities(transcript, config, replacement_source)
            log(f"Found {len(entities)} entities")
            for e in entities:
                log(f"  [{e['start']}:{e['end']}] '{e['original']}' → '{e['replacement']}'")

        job["_transcript"] = transcript

        if batch_mode:
            log("⚠ Batch mode: skipping human review — applying auto-detected entities directly")
            user_entities = entities
        else:
            job["status"] = "review"
            log_q.put({"type": "review", "transcript": transcript, "entities": entities})

            try:
                user_entities = apply_q.get(timeout=1800)
            except queue.Empty:
                job["status"] = "error"
                job["error"] = "Review session expired"
                log_q.put({"type": "error", "msg": "Review session timed out"})
                return

        job["status"] = "running"
        log("▶ Step 3: Synthesizing with " + ("auto-detected entities (batch mode)" if batch_mode else "reviewed entities"))

        with contextlib.redirect_stdout(writer), contextlib.redirect_stderr(writer):
            from audio.pipelines.asr_tts import synthesize_from_entities

            out_path = str(OUTPUTS_DIR / f"{job_id}.wav")
            anonymized_text, _ = synthesize_from_entities(transcript, user_entities, out_path, working_audio, f5_voice)

        if noise_mode == "keep":
            out_path = _mix_noise_if_needed(job, out_path, [], log, writer)

        job["result"] = {
            "transcript": transcript,
            "anonymized": anonymized_text,
            "output_audio": out_path,
        }
        job["status"] = "done"
        log_q.put({"type": "done"})

    except Exception as e:
        job["status"] = "error"
        job["error"] = str(e)
        log_q.put({"type": "log", "msg": traceback.format_exc()})
        log_q.put({"type": "error", "msg": str(e)})
    finally:
        _release_gpu()


def _worker_vc_tts(
    job_id: str,
    audio_path: str,
    f5_voice: str = "impnac_ant",
) -> None:
    """ ASR and re-synthesize with F5-TTS (same content, different voice) """
    job = JOBS[job_id]
    log_q: queue.Queue = job["log_q"]
    writer = _QueueWriter(log_q)
    noise_mode: str = job["noise_mode"]
    device = "cuda" if _cuda_available() else "cpu"

    def log(msg: str) -> None:
        log_q.put({"type": "log", "msg": msg})

    working_audio = audio_path
    transcript = ""
    out_path = ""
    try:
        if noise_mode in ("remove", "keep"):
            working_audio = _separate_if_needed(job, audio_path, log, writer, device)

        with contextlib.redirect_stdout(writer), contextlib.redirect_stderr(writer):
            from audio.asr import get_transcript
            from audio.tts import f5 as tts_f5_module

            log("▶ Step 1: Transcribing audio")
            transcript = get_transcript(working_audio)
            log(f"Transcript: {transcript}")

            log(f"▶ Step 2: Re-synthesizing with voice '{f5_voice}'")
            out_path = str(OUTPUTS_DIR / f"{job_id}.wav")
            tts_f5_module.text_to_speech(transcript, out_path, working_audio, transcript, f5_voice)
            log(f"Saved to {out_path}")

        if noise_mode == "keep":
            out_path = _mix_noise_if_needed(job, out_path, [], log, writer)

        job["result"] = {
            "transcript": None,
            "anonymized": None,
            "output_audio": out_path,
        }
        job["status"] = "done"
        log_q.put({"type": "done"})
    except Exception as e:
        job["status"] = "error"
        job["error"] = str(e)
        log_q.put({"type": "log", "msg": traceback.format_exc()})
        log_q.put({"type": "error", "msg": str(e)})
    finally:
        _release_gpu()


def _worker_vc_knn(
    job_id: str,
    audio_path: str,
    ref_set: str = "",
) -> None:
    """ apply kNN-VC to the full audio (mantains prosody and content) """
    job = JOBS[job_id]
    log_q: queue.Queue = job["log_q"]
    writer = _QueueWriter(log_q)
    noise_mode: str = job["noise_mode"]

    def log(msg: str) -> None:
        log_q.put({"type": "log", "msg": msg})

    out_path = ""
    try:
        with contextlib.redirect_stdout(writer), contextlib.redirect_stderr(writer):
            from audio.pipelines.infilling import apply_knn_vc

            device = "cuda" if _cuda_available() else "cpu"
            log(f"Device: {device}")

        working_audio = audio_path
        if noise_mode in ("remove", "keep"):
            working_audio = _separate_if_needed(job, audio_path, log, writer, device)

        with contextlib.redirect_stdout(writer), contextlib.redirect_stderr(writer):
            log(" applying kNN-VC voice conversion to full audio")
            out_path = str(OUTPUTS_DIR / f"{job_id}.wav")
            vc_tmp = out_path + ".vc_tmp.wav"
            try:
                apply_knn_vc(working_audio, vc_tmp, _knn_ref_wavs(ref_set), device=device)
                os.replace(vc_tmp, out_path)
            finally:
                if os.path.exists(vc_tmp):
                    os.unlink(vc_tmp)
            log(f"Saved to {out_path}")

        if noise_mode == "keep":
            out_path = _mix_noise_if_needed(job, out_path, [], log, writer)

        job["result"] = {
            "transcript": None,
            "anonymized": None,
            "output_audio": out_path,
        }
        job["status"] = "done"
        log_q.put({"type": "done"})
    except Exception as e:
        job["status"] = "error"
        job["error"] = str(e)
        log_q.put({"type": "log", "msg": traceback.format_exc()})
        log_q.put({"type": "error", "msg": str(e)})
    finally:
        _release_gpu()


def _worker_resynth(
    job_id: str,
    audio_path: str,
    words,
    transcript: str,
    entities: list[dict],
    device: str,
    ref_wavs: list[str] | None,
    pipeline: str = "f5",
) -> None:
    """Re-run infilling synthesis with the same entities, no ASR/detection."""
    job = JOBS[job_id]
    log_q: queue.Queue = job["log_q"]
    writer = _QueueWriter(log_q)
    noise_mode: str = job["noise_mode"]

    def log(msg: str) -> None:
        log_q.put({"type": "log", "msg": msg})

    try:
        if pipeline == "f5":
            log("▶ Re-synthesizing with F5 infilling")
            from audio.pipelines.infilling_f5 import synthesize_anonymized
        else:
            log("▶ Re-synthesizing with Google TTS infilling")
            from audio.pipelines.infilling import synthesize_anonymized

        with contextlib.redirect_stdout(writer), contextlib.redirect_stderr(writer):
            out_path = str(OUTPUTS_DIR / f"{job_id}.wav")
            anonymized_text, _, segment_map = synthesize_anonymized(
                audio_path, out_path, words, entities,
                ref_wav_paths=ref_wavs, device=device,
            )

        if noise_mode == "keep":
            out_path = _mix_noise_if_needed(job, out_path, segment_map, log, writer)

        job["_audio_path"] = audio_path
        job["_entities"] = entities
        job["_ref_wavs"] = ref_wavs
        job["_words"] = words
        job["_transcript"] = transcript
        job["result"] = {
            "transcript": transcript,
            "anonymized": anonymized_text,
            "output_audio": out_path,
            "segment_map": segment_map,
        }
        job["status"] = "done"
        log_q.put({"type": "done"})

    except Exception as e:
        job["status"] = "error"
        job["error"] = str(e)
        log_q.put({"type": "log", "msg": traceback.format_exc()})
        log_q.put({"type": "error", "msg": str(e)})
    finally:
        _release_gpu()


# ── HTTP endpoints ────────────────────────────────────────────────────────────

@app.get("/")
def root():
    return FileResponse("static/index.html")


@app.post("/extract-text")
async def extract_text_endpoint(document: UploadFile) -> dict:
    """Extract plain text from an uploaded .txt or .docx file for the Text tab."""
    ext = Path(document.filename or "").suffix.lower()
    data = await document.read()

    doc_id = None
    if ext == ".txt":
        try:
            text = data.decode("utf-8")
        except UnicodeDecodeError:
            text = data.decode("latin-1")
        text = text.strip()
    elif ext == ".docx":
        from docx import Document
        try:
            doc = Document(io.BytesIO(data))
        except Exception:
            raise HTTPException(status_code=400, detail="Could not read .docx file — it may be corrupted.")
        # one line per paragraph, blanks included, so /download-docx can map
        # anonymized lines back onto the original paragraphs 1:1 — table
        # content is intentionally not extracted (v1 scope: paragraphs only)
        text = "\n".join(p.text for p in doc.paragraphs)
        if text.strip():
            doc_id = uuid.uuid4().hex
            (UPLOADS_DIR / f"text_{doc_id}.docx").write_bytes(data)
    else:
        raise HTTPException(status_code=400, detail="Unsupported file type. Please upload a .txt or .docx file.")

    if not text.strip():
        raise HTTPException(status_code=400, detail="No readable text found in the uploaded file.")

    return {"text": text, "doc_id": doc_id}


@app.post("/detect-text")
def detect_text_endpoint(
    text: str = Form(...),
    per: str = Form("rand"),
    loc: str = Form("gen"),
    org: str = Form("gen"),
    misc: str = Form("tag"),
    regex_method: str = Form("tag"),
    whitelist: str = Form(""),
    regex_skip: str = Form(""),
    replacement_source: str = Form("wikidata"),
) -> dict:
    """Detect entities; each entity includes local_match and method.

    replacement_source="wikidata" (default) restricts lookups to the local DB
    only (local_only=True) — a live wikidata.org query is opt-in per entity
    via /lookup-live. replacement_source="llm" instead sources "gen"/"rand"
    replacement text from a locally-hosted Gemma model (never touches the
    network).
    """
    config = _make_config(per, loc, org, misc, regex_method, whitelist, regex_skip)
    try:
        entities = _detect_entities(text, config, replacement_source, local_only=True)
    except RuntimeError as e:
        raise HTTPException(status_code=502, detail=str(e))
    return {"text": text, "entities": entities}


@app.post("/lookup-live")
def lookup_live_endpoint(
    term: str = Form(...),
    ner_label: str = Form(...),
    method: str = Form(...),
) -> dict:
    """Run a live Wikidata query for one entity and return the replacement string."""
    from text_anon import lookup_live
    replacement = lookup_live(term, ner_label, method)
    return {"replacement": replacement}


@app.post("/apply-text")
def apply_text_endpoint(
    text: str = Form(...),
    entities: str = Form(...),
) -> dict:
    """Apply user-reviewed entities to text and return the anonymized result."""
    from text_anon import apply_entities
    ents = json.loads(entities)
    anonymized = apply_entities(text, ents)
    return {"original": text, "anonymized": anonymized}


@app.post("/download-docx")
def download_docx_endpoint(doc_id: str = Form(...), text: str = Form(...)) -> StreamingResponse:
    """Rebuild the originally-uploaded .docx with anonymized text, keeping each
    paragraph's formatting (alignment, indentation, style) — character-level
    formatting (e.g. bold spans) is not preserved, since it's collapsed into
    a single run per paragraph."""
    from docx import Document
    src_path = UPLOADS_DIR / f"text_{doc_id}.docx"
    if not src_path.exists():
        raise HTTPException(status_code=404, detail="Original document not found — please re-upload it.")

    doc = Document(str(src_path))
    lines = text.split("\n")
    for i, p in enumerate(doc.paragraphs):
        new_text = lines[i] if i < len(lines) else ""
        for run in p.runs:
            run.text = ""
        if p.runs:
            p.runs[0].text = new_text
        else:
            p.add_run(new_text)

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": 'attachment; filename="anonymized.docx"'},
    )


@app.post("/jobs")
async def create_job(
    file: UploadFile,
    pipeline: str = Form("infilling"),
    per: str = Form("rand"),
    loc: str = Form("gen"),
    org: str = Form("gen"),
    misc: str = Form("tag"),
    regex_method: str = Form("tag"),
    ref_set: str = Form("al"),
    f5_voice: str = Form("al"),
    noise_mode: str = Form("remove"),
    whitelist: str = Form(""),
    regex_skip: str = Form(""),
    batch_mode: bool = Form(False),
    replacement_source: str = Form("wikidata"),
) -> dict:
    """Accept an audio file, start a pipeline worker, return job_id.

    batch_mode only affects pipelines with a detect→review→synthesize flow
    (infilling / full_infilling_knn / full_asr_tts): when true, the job skips
    the human review/apply pause and synthesizes directly from the
    auto-detected entities. The frontend sets it automatically whenever more
    than one file is submitted at once, since there is no UI to review each
    file individually in that case. It has no effect on vc_tts/vc_knn, which
    never pause for review.

    replacement_source likewise only affects those same three pipelines (the
    ones that call detect_entities) — see _detect_entities().
    """
    if pipeline in ("vc_knn", "full_infilling_knn"):
        try:
            _knn_ref_wavs(ref_set)  # fail before upload/ASR/review, not after
        except RuntimeError as e:
            raise HTTPException(status_code=400, detail=str(e)) from e

    job_id = uuid.uuid4().hex
    raw_path = str(UPLOADS_DIR / f"{job_id}_{file.filename}")
    with open(raw_path, "wb") as f:
        f.write(await file.read())

    audio_path = str(UPLOADS_DIR / f"{job_id}.wav")
    _convert_to_wav16k(raw_path, audio_path)

    JOBS[job_id] = _make_job(noise_mode=noise_mode, batch_mode=batch_mode)

    if pipeline == "infilling":
        target = _worker_detect_infilling
        args = (job_id, audio_path, per, loc, org, misc, regex_method, False, "", whitelist, regex_skip, batch_mode, replacement_source)
    elif pipeline == "vc_tts":
        target = _worker_vc_tts
        args = (job_id, audio_path, f5_voice)
    elif pipeline == "vc_knn":
        target = _worker_vc_knn
        args = (job_id, audio_path, ref_set)
    elif pipeline == "full_asr_tts":
        target = _worker_detect_asr_tts
        args = (job_id, audio_path, per, loc, org, misc, regex_method, f5_voice, whitelist, regex_skip, batch_mode, replacement_source)
    elif pipeline == "full_infilling_knn":
        target = _worker_detect_infilling
        args = (job_id, audio_path, per, loc, org, misc, regex_method, True, ref_set, whitelist, regex_skip, batch_mode, replacement_source)
    else:
        raise HTTPException(status_code=400, detail=f"Unknown pipeline: {pipeline}")

    threading.Thread(target=target, args=args, daemon=True).start()
    return {"job_id": job_id}


@app.post("/jobs/{job_id}/apply")
async def apply_job(job_id: str, entities: str = Form(...)) -> dict:
    """Submit user-reviewed entities to unblock a job waiting in review phase."""
    job = JOBS.get(job_id)
    if not job:
        raise HTTPException(status_code=404, detail="Job not found")
    if job["status"] != "review":
        raise HTTPException(status_code=400, detail=f"Job is not in review state (status={job['status']})")
    ents = json.loads(entities)
    job["apply_q"].put(ents)
    return {"status": "ok"}


@app.post("/jobs/{job_id}/resynth")
async def resynth_job(job_id: str, pipeline: str = "f5") -> dict:
    """Spawn a new synthesis pass reusing the same audio, words and entities."""
    parent = JOBS.get(job_id)
    if not parent:
        raise HTTPException(status_code=404, detail="Job not found")
    missing = [k for k in ("_audio_path", "_words", "_entities") if k not in parent]
    if missing:
        raise HTTPException(status_code=400, detail=f"Job missing re-synthesis data: {missing}")

    new_job_id = uuid.uuid4().hex
    JOBS[new_job_id] = _make_job(noise_mode=parent.get("noise_mode", "remove"))
    JOBS[new_job_id]["_noise_path"] = parent.get("_noise_path")
    JOBS[new_job_id]["_original_audio_path"] = parent.get("_original_audio_path")
    device = "cuda" if _cuda_available() else "cpu"
    threading.Thread(
        target=_worker_resynth,
        args=(
            new_job_id,
            parent["_audio_path"],
            parent["_words"],
            parent.get("_transcript", ""),
            parent["_entities"],
            device,
            parent.get("_ref_wavs"),
            pipeline,
        ),
        daemon=True,
    ).start()
    return {"job_id": new_job_id}


@app.get("/jobs/{job_id}/stream")
def stream_job(job_id: str) -> StreamingResponse:
    """Pull messages from queue and stream as SSE for live updates."""
    job = JOBS.get(job_id)
    if not job:
        raise HTTPException(status_code=404, detail="Job not found")

    def generator():
        log_q: queue.Queue = job["log_q"]
        while True:
            try:
                msg = log_q.get(timeout=30)
                yield f"data: {json.dumps(msg)}\n\n"
                if msg["type"] in ("done", "error"):
                    break
            except queue.Empty:
                yield 'data: {"type":"ping"}\n\n'

    return StreamingResponse(
        generator(),
        media_type="text/event-stream",
        headers={"Cache-Control": "no-cache", "X-Accel-Buffering": "no"},
    )


@app.get("/jobs/{job_id}/result")
def get_result(job_id: str) -> dict:
    """Return job's final output."""
    job = JOBS.get(job_id)
    if not job:
        raise HTTPException(status_code=404, detail="Job not found")
    return {
        "status": job["status"],
        "result": job["result"],
        "error": job["error"],
    }


@app.get("/jobs/{job_id}/audio")
def get_audio(job_id: str) -> FileResponse:
    """Return the output .wav file once processing is done."""
    job = JOBS.get(job_id)
    if not job or not job.get("result"):
        raise HTTPException(status_code=404, detail="Audio not ready")
    out_path = job["result"]["output_audio"]
    if not out_path or not os.path.exists(out_path):
        raise HTTPException(status_code=404, detail="Audio file not found")
    return FileResponse(out_path, media_type="audio/wav", filename="anonymized.wav")


@app.get("/jobs/{job_id}/segment_map")
def get_segment_map(job_id: str):
    """Return the segment map as a downloadable JSON file."""
    job = JOBS.get(job_id)
    if not job or not job.get("result"):
        raise HTTPException(status_code=404, detail="Job not ready")
    seg_map = job["result"].get("segment_map")
    if seg_map is None:
        raise HTTPException(status_code=404, detail="No segment map for this job")
    from fastapi.responses import Response
    return Response(
        content=json.dumps(seg_map, indent=2),
        media_type="application/json",
        headers={"Content-Disposition": f'attachment; filename="segment_map_{job_id}.json"'},
    )


@app.get("/f5_preview/{ref_set}")
def f5_preview(ref_set: str) -> FileResponse:
    """Serve the 12s F5-TTS reference clip for the given speaker (generated by make_f5_refs.py)."""
    from audio.tts.f5 import F5_REFS_DIR
    p = Path(F5_REFS_DIR) / f"{ref_set}.wav"
    if not p.exists():
        raise HTTPException(status_code=404, detail=f"F5 ref not found for '{ref_set}' — run make_f5_refs.py")
    return FileResponse(str(p), media_type="audio/wav")


@app.get("/ref_preview/{ref_set}")
def ref_preview(ref_set: str) -> FileResponse:
    """Serve the pre-generated voice-converted preview for the given speaker."""
    p = Path(__file__).parent / "ref" / "previews" / f"{ref_set}.wav"
    if not p.exists():
        raise HTTPException(status_code=404, detail=f"Preview not found for '{ref_set}' — run make_ref_previews.py")
    return FileResponse(str(p), media_type="audio/wav")


# ── Static files & entry point ────────────────────────────────────────────────

app.mount("/static", StaticFiles(directory=str(APP_DIR / "static")), name="static")

if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Speech Anonymizer web platform")
    parser.add_argument("--host", default="127.0.0.1", help="Bind host (default: 127.0.0.1)")
    parser.add_argument("--port", type=int, default=8000, help="Port (default: 8000)")
    args = parser.parse_args()
    print(f"Starting at http://{args.host}:{args.port}")
    uvicorn.run(app, host=args.host, port=args.port)
